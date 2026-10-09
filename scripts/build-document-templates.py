"""Build editable document downloads from the same locale content as the library."""
from pathlib import Path
import argparse
import json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'assets/document-templates'


def base_document(locale, contract=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    margin = 25 if contract else 18
    sec.top_margin = sec.bottom_margin = Mm(margin)
    sec.left_margin = sec.right_margin = Mm(28 if contract else 20)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        fonts = style._element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if 'Theme' in attribute:
                del fonts.attrib[attribute]
        fonts.set(qn('w:eastAsia'), 'KaiTi')
        fonts.set(qn('w:cs'), 'Times New Roman')
        language = OxmlElement('w:lang')
        language.set(qn('w:val'), 'en-US')
        language.set(qn('w:eastAsia'), 'zh-TW')
        style._element.rPr.append(language)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal'].paragraph_format
    doc.styles['Normal'].font.size = Pt(11 if contract else 10.5)
    normal.line_spacing = 1.25 if contract else 1.12
    normal.space_after = Pt(6 if contract else 3)
    normal.widow_control = True
    doc.styles['Title'].font.size = Pt(14 if contract else 18)
    doc.styles['Title'].font.bold = contract
    doc.styles['Title'].paragraph_format.space_after = Pt(22 if contract else 8)
    heading = doc.styles['Heading 1']
    heading.font.size = Pt(11 if contract else 12)
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(12 if contract else 7)
    heading.paragraph_format.space_after = Pt(6 if contract else 3)
    heading.paragraph_format.keep_with_next = True
    # Explicit script fonts avoid a Japanese theme fallback in Word.
    for element in doc.styles.element.iter(qn('w:rFonts')):
        for attribute in list(element.attrib):
            if 'Theme' in attribute:
                del element.attrib[attribute]
        for script in ['ascii', 'hAnsi', 'cs']:
            element.set(qn('w:' + script), 'Times New Roman')
        element.set(qn('w:eastAsia'), 'KaiTi')
    # Use the installed KaiTi family instead of relying on an unavailable font.
    font_table = doc.part.package.part_related_by('http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument').related_parts
    for part in font_table.values():
        if str(part.partname).endswith('/fontTable.xml'):
            from lxml import etree
            tree = etree.fromstring(part.blob)
            font = OxmlElement('w:font'); font.set(qn('w:name'), 'KaiTi')
            alternate = OxmlElement('w:altName'); alternate.set(qn('w:val'), 'KaiTi'); font.append(alternate)
            tree.append(font); part._blob = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
    doc.core_properties.author = 'Taloox'
    doc.core_properties.comments = ''
    return doc


def numbered_items(doc, text):
    """Start a native Word list for each article, with hanging indentation."""
    numbering = doc.part.numbering_part.element
    ids = [int(e.get(qn('w:abstractNumId'))) for e in numbering.findall(qn('w:abstractNum'))]
    abstract_id = max(ids, default=-1) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
        element = OxmlElement('w:' + tag); element.set(qn('w:val'), value); level.append(element)
    abstract.append(level); numbering.append(abstract)
    num_id = numbering.add_num(abstract_id).numId
    for line in text.split('\n'):
        p = doc.add_paragraph(line)
        fmt = p.paragraph_format
        fmt.left_indent, fmt.first_line_indent = Mm(10), Mm(-6)
        fmt.keep_together = True
        properties = p._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_ilvl().val = 0
        properties.get_or_add_numId().val = num_id


def article(doc, common, locale, number, heading, body):
    label = common['numberLabels'].split('|')[number - 1]
    p = doc.add_paragraph(common['article'].replace('{number}', label) + ' ' + heading, 'Heading 1')
    if number == 5:
        p.paragraph_format.page_break_before = True
    numbered_items(doc, body)


def build_contract(copy, item, locale):
    doc = base_document(locale, contract=True)
    c, data = copy['contractCommon'], item['contract']
    doc.core_properties.title = data['title']
    doc.core_properties.subject = 'Editable agreement template'
    title = doc.add_paragraph(data['title'], 'Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(c['partiesTitle']); p.paragraph_format.space_after = Pt(4)
    for role in ['partyA', 'partyB']:
        p = doc.add_paragraph(f"{data[role]} {c['legalName']} ({c['partySuffix']} {c[role]})")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph(c['intro']); p.paragraph_format.space_before = Pt(14)
    article(doc, c, locale, 1, c['scopeTitle'], item['scope'])
    article(doc, c, locale, 2, c['termTitle'], data['term'])
    article(doc, c, locale, 3, c['acceptanceTitle'], data['acceptance'])
    # Balance the terms across two pages and keep execution on its own page.
    article(doc, c, locale, 4, c['feesTitle'], c['fees'])
    article(doc, c, locale, 5, c['paymentTitle'], c['payment'])
    p = doc.add_paragraph(c['bank']); p.paragraph_format.left_indent = Mm(10)
    article(doc, c, locale, 6, data['rightsTitle'], item['rights'])
    article(doc, c, locale, 7, c['cancelTitle'], copy['common']['cancel'])
    article(doc, c, locale, 8, c['legalTitle'], c['legal'])
    article(doc, c, locale, 9, c['otherTitle'], c['other'])
    p = doc.add_paragraph(c['signTitle'], 'Heading 1')
    p.paragraph_format.page_break_before = True
    for role in ['partyA', 'partyB']:
        p = doc.add_paragraph(f"{c[role]}  {data[role]}")
        p.paragraph_format.space_before = Pt(16)
        p.runs[0].bold = True
        for line in c['signFields'].split('\n'):
            doc.add_paragraph(line)
        p = doc.add_paragraph(c['signature'] + '  ______________________________')
        p.paragraph_format.space_before = Pt(14)
        doc.add_paragraph(c['date'])
    p = doc.add_paragraph(c['executionDate']); p.paragraph_format.space_before = Pt(16)
    return doc


def build_quote(copy, item, locale):
    doc = base_document(locale)
    c = copy['common']
    doc.core_properties.title = f"{item['title']} {copy['quote']}"
    doc.core_properties.subject = 'Editable quotation template'
    doc.add_paragraph(doc.core_properties.title, 'Title')
    doc.add_paragraph(c['quoteIntro'])
    def section(heading, body):
        doc.add_paragraph(heading, 'Heading 1'); doc.add_paragraph(body)
    section(c['partiesTitle'], c['parties'])
    section(c['scopeTitle'], item['scope'])
    doc.add_paragraph(c['lineTitle'], 'Heading 1'); doc.add_paragraph(c['lineHelp'])
    table = doc.add_table(rows=1, cols=4); table.style = 'Light Shading Accent 1'
    for cell, text in zip(table.rows[0].cells, [copy['table'][key] for key in ['item', 'quantity', 'unitPrice', 'amount']]):
        cell.text = text
    for name in item['lines'].split(' | '):
        for cell, text in zip(table.add_row().cells, [name, '[___]', '[___]', '[___]']):
            cell.text = text
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, [3.2, 1.1, 1, 1]):
            cell.width = Inches(width)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    section(c['paymentTitle'], c['totals'])
    section(c['rightsTitle'], item['rights'])
    section(c['quoteTermsTitle'], c['quoteTerms'])
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--kind', choices=['contract', 'quote', 'all'], default='all')
    args = parser.parse_args()
    OUTPUT.mkdir(exist_ok=True)
    count = 0
    for locale in ['en', 'zh-TW']:
        copy = json.loads((ROOT / f'messages/{locale}.json').read_text())['documentLibrary']
        for key, item in copy['documents'].items():
            for kind, build in [('contract', build_contract), ('quote', build_quote)]:
                if args.kind not in ['all', kind]:
                    continue
                doc = build(copy, item, locale)
                for paragraph in list(doc.paragraphs) + [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]:
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'KaiTi')
                doc.save(OUTPUT / f'taloox-{key}-{kind}-{locale}.docx')
                count += 1
    print(f'Updated {count} editable templates')


if __name__ == '__main__':
    main()
